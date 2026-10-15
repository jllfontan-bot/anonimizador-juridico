import sys
import tempfile
import unittest
from pathlib import Path

import cv2
import docx
import numpy as np
from odf.opendocument import OpenDocumentText, load as cargar_odt
from odf.style import Style, TextProperties
from odf.text import P, Span
from odf.draw import Frame, Image as ImaxeOdt

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from anonimizador import web
from anonimizador.pipeline import (
    anonimizar_e_gardar,
    anonimizar_texto,
    cargar_documento_estruturado,
    gardar_documento_estruturado,
    ruta_saida_por_defecto,
)
from anonimizador.registry import RexistroCorrespondencias


def _png_qr(contido: str, tamano_modulo: int = 8, marxe_modulos: int = 4) -> bytes:
    """Xera un QR real coma PNG usando SÓ `cv2.QRCodeEncoder` (sen
    depender da libraría `qrcode`, que non fai falta en produción)."""
    modulos = cv2.QRCodeEncoder.create().encode(contido)
    con_marxe = np.pad(modulos, marxe_modulos, constant_values=255)
    grande = np.kron(con_marxe, np.ones((tamano_modulo, tamano_modulo), dtype=np.uint8))
    return cv2.imencode(".png", grande)[1].tobytes()


def _ten_qr_detectable(datos_png: bytes) -> bool:
    arr = np.frombuffer(datos_png, dtype=np.uint8)
    imaxe = cv2.imdecode(arr, cv2.IMREAD_GRAYSCALE)
    return bool(cv2.QRCodeDetector().detectMulti(imaxe)[0])


def _crear_docx_con_formato(ruta: Path) -> None:
    doc = docx.Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Manuel Ferrándiz Salcedo ")
    r1.bold = True
    r2 = p.add_run("declarou perante o xulgado.")
    r2.italic = True

    tabla = doc.add_table(rows=1, cols=2)
    tabla.cell(0, 0).paragraphs[0].add_run("DNI:").bold = True
    tabla.cell(0, 1).paragraphs[0].add_run("30014811Z")

    doc.save(str(ruta))


def _crear_odt_con_formato(ruta: Path) -> None:
    doc = OpenDocumentText()
    estilo = Style(name="Negra", family="text")
    estilo.addElement(TextProperties(fontweight="bold"))
    doc.automaticstyles.addElement(estilo)

    p = P()
    p.addElement(Span(stylename=estilo, text="Manuel Ferrándiz Salcedo"))
    p.addText(" declarou perante o xulgado.")
    doc.text.addElement(p)
    doc.save(str(ruta))


def _crear_docx_nome_partido_en_dous_runs(ruta: Path) -> None:
    """O nome da persoa queda repartido entre dúas runs con formato
    distinto -- o caso máis esixente de mestura dentro dun parágrafo."""
    doc = docx.Document()
    p = doc.add_paragraph()
    p.add_run("Manuel ")  # sen formato
    r2 = p.add_run("Ferrándiz Salcedo")
    r2.bold = True
    p.add_run(" declarou perante o xulgado.")
    doc.save(str(ruta))


class TestFormatoDocx(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.rexistro = RexistroCorrespondencias(str(self.dir / "rexistro"))

    def tearDown(self):
        self._tmp.cleanup()

    def test_conserva_formato_do_primeiro_run_e_paragrafo(self):
        entrada = self.dir / "escrito.docx"
        _crear_docx_con_formato(entrada)
        saida = self.dir / "escrito.anonimizado.docx"

        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        doc = docx.Document(str(saida))
        p = doc.paragraphs[0]
        self.assertIn("PERSOA_000001", p.text)
        self.assertNotIn("Manuel Ferrándiz Salcedo", p.text)
        self.assertTrue(p.runs[0].bold)  # conserva o formato da primeira run

    def test_conserva_formato_de_TODAS_as_runs_non_so_a_primeira(self):
        """A run en cursiva, sen tocar pola substitución, ten que quedar
        exactamente igual -- non se pode colapsar todo no formato da
        primeira run."""
        entrada = self.dir / "escrito.docx"
        _crear_docx_con_formato(entrada)
        saida = self.dir / "escrito.anonimizado.docx"

        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        doc = docx.Document(str(saida))
        p = doc.paragraphs[0]
        self.assertEqual(len(p.runs), 2)
        self.assertTrue(p.runs[0].bold)
        self.assertIn("PERSOA_000001", p.runs[0].text)
        self.assertEqual(p.runs[1].text, "declarou perante o xulgado.")
        self.assertTrue(p.runs[1].italic)

    def test_substitucion_que_atravesa_dous_runs_non_se_duplica(self):
        """O nome vai partido entre unha run sen formato e outra en
        negriña; a mención cóbrea NER coma unha soa entidade. O código
        debe aparecer UNHA soa vez, herdando o formato da run onde empeza."""
        entrada = self.dir / "escrito.docx"
        _crear_docx_nome_partido_en_dous_runs(entrada)
        saida = self.dir / "escrito.anonimizado.docx"

        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        doc = docx.Document(str(saida))
        p = doc.paragraphs[0]
        self.assertEqual(p.text.count("PERSOA_000001"), 1)
        self.assertNotIn("Manuel", p.text)
        self.assertNotIn("Ferrándiz Salcedo", p.text)
        self.assertIn("declarou perante o xulgado.", p.text)
        # a run 0 (sen formato orixinal) e a que recibe o código
        self.assertIn("PERSOA_000001", p.runs[0].text)
        self.assertFalse(p.runs[0].bold)
        # a run 1 (orixinalmente en negriña) queda baleira, non duplicada
        self.assertEqual(p.runs[1].text, "")

    def test_conserva_estrutura_de_taboa(self):
        entrada = self.dir / "escrito.docx"
        _crear_docx_con_formato(entrada)
        saida = self.dir / "escrito.anonimizado.docx"

        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        doc = docx.Document(str(saida))
        self.assertEqual(len(doc.tables), 1)
        self.assertEqual(doc.tables[0].cell(0, 0).text, "DNI:")
        self.assertNotIn("30014811Z", doc.tables[0].cell(0, 1).text)

    def test_saida_por_defecto_mantén_extension_docx(self):
        self.assertTrue(ruta_saida_por_defecto("/x/escrito.docx").endswith(".anonimizado.docx"))

    def test_paragrafo_con_salto_de_liña_manual_conserva_o_formato(self):
        """Caso real: unha carta cun enderezo escrito coma UN SÓ parágrafo
        con Shift+Enter no medio ("Rúa Maior 5" salto "Vigo, 36001") --
        python-docx representa ese salto manual coma un "\\n" DENTRO do
        texto do propio parágrafo. Tratar "1 liña = 1 parágrafo" (como
        antes) facía que a conta de liñas nunca coincidise coa de
        parágrafos, e o documento caía sempre no respaldo de texto plano
        -- perdendo o .docx orixinal -- aínda SEN editar nada a man."""
        entrada = self.dir / "carta.docx"
        doc = docx.Document()
        p1 = doc.add_paragraph()
        p1.add_run("Manuel Ferrándiz Salcedo")
        p1.add_run().add_break()
        p1.add_run("Rúa Maior 5, Vigo")
        doc.add_paragraph("Un saúdo,")
        doc.save(str(entrada))

        saida = self.dir / "carta.anonimizado.docx"
        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        self.assertEqual(saida.suffix, ".docx")
        doc_final = docx.Document(str(saida))
        self.assertNotIn("Manuel Ferrándiz Salcedo", doc_final.paragraphs[0].text)
        self.assertIn("PERSOA_000001", doc_final.paragraphs[0].text)
        self.assertNotIn("Rúa Maior 5", doc_final.paragraphs[0].text)
        self.assertIn("[ENDEREZO]", doc_final.paragraphs[0].text)

    def test_liñas_editadas_de_menos_falla_con_erro_claro(self):
        """Se se edita o texto na revisión e se borra un salto de liña, non
        se debe escribir un documento desaliñado -- debe fallar."""
        entrada = self.dir / "escrito.docx"
        _crear_docx_con_formato(entrada)
        texto, contexto = cargar_documento_estruturado(str(entrada))
        texto_anon, _ = anonimizar_texto(texto, self.rexistro)
        texto_editado = texto_anon.replace("\n", " ")  # fusiona dúas liñas nunha
        with self.assertRaises(ValueError):
            gardar_documento_estruturado(contexto, texto_editado, str(self.dir / "saida.docx"))

    def test_qr_incrustado_no_docx_queda_redactado(self):
        """Un QR de verificación incrustado coma IMAXE (non forma parte de
        ningún parágrafo nin substitución de texto) tamén ten que
        redactarse -- mesmo principio ca en PDF."""
        entrada = self.dir / "carta_con_qr.docx"
        doc = docx.Document()
        doc.add_paragraph("Manuel Ferrándiz Salcedo presenta este escrito.")
        doc.add_paragraph().add_run().add_picture(
            __import__("io").BytesIO(_png_qr("CSV:VERIFICACION-REAL")), width=docx.shared.Inches(1)
        )
        doc.save(str(entrada))

        saida = self.dir / "carta_con_qr.anonimizado.docx"
        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        self.assertEqual(saida.suffix, ".docx")
        doc_final = docx.Document(str(saida))
        imaxes = [
            p for p in doc_final.part.related_parts.values()
            if getattr(p, "content_type", "").startswith("image/")
        ]
        self.assertEqual(len(imaxes), 1)
        self.assertFalse(_ten_qr_detectable(imaxes[0].blob))

    def test_qr_na_cabeceira_tamen_queda_redactado(self):
        """Caso real atopado arranxando isto: unha imaxe incrustada na
        CABECEIRA (non no corpo do documento) non aparece en
        `document.part.related_parts` (só nas relacións da súa propia
        HeaderPart) -- hai que percorrer TODO o paquete (iter_parts()),
        non só as relacións directas do documento principal."""
        entrada = self.dir / "carta_con_qr_cabeceira.docx"
        doc = docx.Document()
        doc.add_paragraph("Manuel Ferrándiz Salcedo presenta este escrito.")
        cabeceira = doc.sections[0].header
        cabeceira.paragraphs[0].add_run().add_picture(
            __import__("io").BytesIO(_png_qr("CSV:CABECEIRA-REAL")), width=docx.shared.Inches(1)
        )
        doc.save(str(entrada))

        saida = self.dir / "carta_con_qr_cabeceira.anonimizado.docx"
        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        doc_final = docx.Document(str(saida))
        imaxes = [
            p for p in doc_final.part.package.iter_parts()
            if getattr(p, "content_type", "").startswith("image/")
        ]
        self.assertTrue(imaxes)
        self.assertFalse(any(_ten_qr_detectable(p.blob) for p in imaxes))

    def test_imaxe_sen_qr_non_se_altera(self):
        """Non debe haber falsos positivos: unha imaxe normal (sen QR nin
        código de barras, ex. un logotipo) ten que quedar intacta."""
        entrada = self.dir / "carta_con_logo.docx"
        doc = docx.Document()
        doc.add_paragraph("Manuel Ferrándiz Salcedo presenta este escrito.")
        logo = np.full((100, 100), 200, dtype=np.uint8)
        _, logo_png = cv2.imencode(".png", logo)
        doc.add_paragraph().add_run().add_picture(
            __import__("io").BytesIO(logo_png.tobytes()), width=docx.shared.Inches(1)
        )
        doc.save(str(entrada))

        saida = self.dir / "carta_con_logo.anonimizado.docx"
        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        doc_final = docx.Document(str(saida))
        imaxes = [
            p for p in doc_final.part.related_parts.values()
            if getattr(p, "content_type", "").startswith("image/")
        ]
        self.assertEqual(len(imaxes), 1)
        self.assertEqual(imaxes[0].blob, logo_png.tobytes())


class TestFormatoOdt(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.rexistro = RexistroCorrespondencias(str(self.dir / "rexistro"))

    def tearDown(self):
        self._tmp.cleanup()

    def test_conserva_texto_e_documento_valido(self):
        entrada = self.dir / "escrito.odt"
        _crear_odt_con_formato(entrada)
        saida = self.dir / "escrito.anonimizado.odt"

        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        doc = cargar_odt(str(saida))
        from odf import teletype
        p = doc.getElementsByType(P)[0]
        texto = teletype.extractText(p)
        self.assertIn("PERSOA_000001", texto)
        self.assertNotIn("Manuel Ferrándiz Salcedo", texto)

    def test_saida_por_defecto_mantén_extension_odt(self):
        self.assertTrue(ruta_saida_por_defecto("/x/escrito.odt").endswith(".anonimizado.odt"))

    def test_conserva_estilo_do_span_e_texto_seguinte_sen_tocar(self):
        """O "span" en negriña co nome ten que quedar co código pero seguir
        sendo un span co MESMO estilo; o texto plano despois del (sen
        formato) ten que quedar exactamente igual, sen fusionarse nun só
        bloque."""
        from odf import teletype
        from odf.element import Text as TextoOdt

        entrada = self.dir / "escrito.odt"
        _crear_odt_con_formato(entrada)
        saida = self.dir / "escrito.anonimizado.odt"

        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        doc = cargar_odt(str(saida))
        p = doc.getElementsByType(P)[0]
        fillos = list(p.childNodes)
        self.assertEqual(len(fillos), 2)  # segue habendo un span + un texto solto

        span, texto_solto = fillos
        self.assertNotIsInstance(span, TextoOdt)
        self.assertEqual(span.getAttribute("stylename"), "Negra")
        self.assertIn("PERSOA_000001", teletype.extractText(span))
        self.assertNotIn("Manuel Ferrándiz Salcedo", teletype.extractText(span))

        self.assertIsInstance(texto_solto, TextoOdt)
        self.assertEqual(texto_solto.data, " declarou perante o xulgado.")

    def test_qr_incrustado_no_odt_queda_redactado(self):
        entrada = self.dir / "carta_con_qr.odt"
        doc = OpenDocumentText()
        doc.text.addElement(P(text="Manuel Ferrándiz Salcedo presenta este escrito."))
        href = doc.addPictureFromString(_png_qr("CSV:VERIFICACION-REAL"), "image/png")
        frame = Frame(width="3cm", height="3cm", anchortype="paragraph")
        frame.addElement(ImaxeOdt(href=href))
        p2 = P()
        p2.addElement(frame)
        doc.text.addElement(p2)
        doc.save(str(entrada))

        saida = self.dir / "carta_con_qr.anonimizado.odt"
        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        self.assertEqual(saida.suffix, ".odt")
        doc_final = cargar_odt(str(saida))
        self.assertEqual(len(doc_final.Pictures), 1)
        (_, datos_imaxe, _) = list(doc_final.Pictures.values())[0]
        self.assertFalse(_ten_qr_detectable(datos_imaxe))


class TestFormatoWeb(unittest.TestCase):
    def setUp(self):
        web.SESIONS.clear()
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.client = web.app.test_client()

    def tearDown(self):
        self._tmp.cleanup()

    def test_fluxo_web_conserva_formato_docx(self):
        entrada = self.dir / "escrito.docx"
        _crear_docx_con_formato(entrada)
        rexistro_dir = str(self.dir / "rexistro")

        resposta = self.client.post(
            "/procesar", data={"ruta": str(entrada), "rexistro": rexistro_dir}, follow_redirects=True
        )
        sid = resposta.request.path.rsplit("/", 1)[-1]
        sesion = web.SESIONS[sid]

        resposta2 = self.client.post(
            f"/revisar/{sid}/confirmar", data={"texto_final": sesion["texto_anonimizado"]}
        )
        self.assertIn("Revisión confirmada", resposta2.get_data(as_text=True))

        saida_esperada = str(entrada.with_suffix("")) + ".anonimizado.docx"
        self.assertTrue(Path(saida_esperada).exists())
        doc = docx.Document(saida_esperada)
        self.assertIn("PERSOA_000001", doc.paragraphs[0].text)
        self.assertTrue(doc.paragraphs[0].runs[0].bold)

    def test_fusion_de_homonimo_conserva_formato_por_tramos(self):
        """Tras confirmar unha fusión, o código actualizado no texto ten
        que seguir escribíndose respectando o formato do run orixinal, non
        caer á reserva de "todo o parágrafo dun só formato"."""
        doc1 = self.dir / "doc1.docx"
        _crear_docx_con_formato(doc1)
        doc2 = self.dir / "doc2.docx"
        _crear_docx_con_formato(doc2)
        rexistro_dir = str(self.dir / "rexistro")

        r1 = self.client.post("/procesar", data={"ruta": str(doc1), "rexistro": rexistro_dir}, follow_redirects=True)
        sid1 = r1.request.path.rsplit("/", 1)[-1]
        self.client.post(f"/revisar/{sid1}/confirmar", data={"texto_final": web.SESIONS[sid1]["texto_anonimizado"]})

        r2 = self.client.post("/procesar", data={"ruta": str(doc2), "rexistro": rexistro_dir}, follow_redirects=True)
        sid2 = r2.request.path.rsplit("/", 1)[-1]
        sesion2 = web.SESIONS[sid2]

        self.client.post(
            f"/revisar/{sid2}/confirmar",
            data={
                "texto_final": sesion2["texto_anonimizado"],
                "decision_PERSOA_000002": "fusionar:PERSOA_000001",
            },
        )

        saida = str(doc2.with_suffix("")) + ".anonimizado.docx"
        doc = docx.Document(saida)
        p = doc.paragraphs[0]
        self.assertEqual(len(p.runs), 2)  # non caeu á reserva (senón serían menos/máis runs)
        self.assertIn("PERSOA_000001", p.runs[0].text)
        self.assertTrue(p.runs[0].bold)
        self.assertEqual(p.runs[1].text, "declarou perante o xulgado.")
        self.assertTrue(p.runs[1].italic)

    def test_edicion_que_desaliña_paragrafos_conserva_o_docx_e_garda_a_edicion_aparte(self):
        """Antes caía a .txt (perdíase o formato). Agora o .docx gárdase coas
        decisións e a edición libre queda nun .edicion-manual.txt (04/10/2026)."""
        entrada = self.dir / "escrito.docx"
        _crear_docx_con_formato(entrada)
        rexistro_dir = str(self.dir / "rexistro")

        resposta = self.client.post(
            "/procesar", data={"ruta": str(entrada), "rexistro": rexistro_dir}, follow_redirects=True
        )
        sid = resposta.request.path.rsplit("/", 1)[-1]
        sesion = web.SESIONS[sid]
        texto_desaliñado = sesion["texto_anonimizado"].replace("\n", " ")

        resposta2 = self.client.post(
            f"/revisar/{sid}/confirmar", data={"texto_final": texto_desaliñado}
        )
        texto = resposta2.get_data(as_text=True)
        self.assertIn("edición libre", texto.lower())

        saida = Path(web.SESIONS[sid]["ruta_saida"])
        self.assertEqual(saida.suffix, ".docx")
        self.assertTrue(saida.with_suffix(".edicion-manual.txt").exists())
        self.assertFalse(Path(str(entrada.with_suffix("")) + ".anonimizado.txt").exists())


if __name__ == "__main__":
    unittest.main()
