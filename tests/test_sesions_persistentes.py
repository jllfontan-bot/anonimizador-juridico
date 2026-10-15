"""Persistencia das sesións de revisión pendentes entre reinicios do
servidor web (`web._obter_sesion`, `registry.gardar_sesion`).

Punto crítico a verificar (informe, sección 03: "quen a lea reverte todo o
sistema de golpe"): o ficheiro `sesions.enc` ten que estar cifrado coma
calquera outro dato do rexistro, e NUNCA debe conter en claro o documento
orixinal sen anonimizar (`contexto`) -- iso recárgase sempre desde o
ficheiro orixinal no disco, nunca se persiste.
"""

import sys
import tempfile
import unittest
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from anonimizador import web
from anonimizador.registry import RexistroCorrespondencias


def _crear_docx_con_formato(ruta: Path) -> None:
    doc = docx.Document()
    p = doc.add_paragraph()
    r1 = p.add_run("Manuel Ferrándiz Salcedo ")
    r1.bold = True
    r2 = p.add_run("declarou perante o xulgado.")
    r2.italic = True
    doc.save(str(ruta))


class TestSesionsPersistentesEnDisco(unittest.TestCase):
    def setUp(self):
        web.SESIONS.clear()
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.rexistro_dir = str(self.dir / "rexistro")
        self.client = web.app.test_client()

    def tearDown(self):
        self._tmp.cleanup()

    def _procesar(self, ruta: str) -> str:
        resposta = self.client.post(
            "/procesar", data={"ruta": ruta, "rexistro": self.rexistro_dir}, follow_redirects=True
        )
        self.assertEqual(resposta.status_code, 200)
        return resposta.request.path.rsplit("/", 1)[-1]

    def test_sesion_pendente_queda_cifrada_en_disco(self):
        doc = self.dir / "declaracion.txt"
        doc.write_text("Manuel Ferrándiz Salcedo, con DNI 30014811-Z, declarou.", encoding="utf-8")

        self._procesar(str(doc))

        ficheiro_sesions = Path(self.rexistro_dir) / "sesions.enc"
        self.assertTrue(ficheiro_sesions.exists())
        contido_binario = ficheiro_sesions.read_bytes()
        # o nome e o DNI reais (detectados e gardados no informe da sesión
        # para que o letrado os revise) NUNCA deben aparecer en claro --
        # teñen que quedar dentro do cifrado Fernet, coma o resto do rexistro
        self.assertNotIn(b"Manuel Ferrandiz Salcedo", contido_binario)
        self.assertNotIn("Manuel Ferrándiz Salcedo".encode("utf-8"), contido_binario)
        self.assertNotIn(b"30014811", contido_binario)

    def test_reconstrue_sesion_tras_reinicio_do_servidor(self):
        doc = self.dir / "declaracion.txt"
        doc.write_text("Manuel Ferrándiz Salcedo, con DNI 30014811-Z, declarou.", encoding="utf-8")

        sid = self._procesar(str(doc))

        # simula un reinicio do servidor: perdese todo o estado en memoria
        web.SESIONS.clear()

        resposta = self.client.get(f"/revisar/{sid}?rexistro={self.rexistro_dir}")
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("PERSOA_000001", resposta.get_data(as_text=True))

    def test_reconstrue_e_confirma_docx_conservando_formato_tras_reinicio(self):
        ruta_docx = self.dir / "escrito.docx"
        _crear_docx_con_formato(ruta_docx)

        sid = self._procesar(str(ruta_docx))
        sesion_orixinal = web.SESIONS[sid]
        texto_anonimizado = sesion_orixinal["texto_anonimizado"]

        # simula un reinicio do servidor xusto antes de confirmar
        web.SESIONS.clear()

        pax = self.client.get(f"/revisar/{sid}?rexistro={self.rexistro_dir}")
        self.assertEqual(pax.status_code, 200)
        # o formulario de confirmación leva o cartafol do rexistro nun campo
        # oculto, para que a ruta de recuperación funcione tamén por POST
        self.assertIn(f'value="{self.rexistro_dir}"', pax.get_data(as_text=True))

        web.SESIONS.clear()  # simula que a consulta anterior foi noutro proceso
        resposta = self.client.post(
            f"/revisar/{sid}/confirmar",
            data={"texto_final": texto_anonimizado, "rexistro_dir": self.rexistro_dir},
        )
        self.assertEqual(resposta.status_code, 200)
        self.assertIn("Revisión confirmada", resposta.get_data(as_text=True))

        ruta_saida = ruta_docx.with_suffix("").as_posix() + ".anonimizado.docx"
        doc_final = docx.Document(ruta_saida)
        self.assertIn("PERSOA_000001", doc_final.paragraphs[0].text)
        self.assertTrue(doc_final.paragraphs[0].runs[0].bold)

    def test_sesion_confirmada_desaparece_do_disco(self):
        doc = self.dir / "declaracion.txt"
        doc.write_text("Manuel Ferrándiz Salcedo, con DNI 30014811-Z, declarou.", encoding="utf-8")

        sid = self._procesar(str(doc))
        rexistro = RexistroCorrespondencias(self.rexistro_dir)
        self.assertIn(sid, rexistro.ler_sesions_pendentes())

        sesion = web.SESIONS[sid]
        self.client.post(
            f"/revisar/{sid}/confirmar",
            data={"texto_final": sesion["texto_anonimizado"], "rexistro_dir": self.rexistro_dir},
        )

        rexistro_recargado = RexistroCorrespondencias(self.rexistro_dir)
        self.assertNotIn(sid, rexistro_recargado.ler_sesions_pendentes())

    def test_index_amosa_pendentes_dun_rexistro_tras_reinicio(self):
        doc = self.dir / "declaracion.txt"
        doc.write_text("Manuel Ferrándiz Salcedo, con DNI 30014811-Z, declarou.", encoding="utf-8")

        self._procesar(str(doc))
        web.SESIONS.clear()

        resposta = self.client.get(f"/?rexistro={self.rexistro_dir}")
        self.assertIn(str(doc), resposta.get_data(as_text=True))


if __name__ == "__main__":
    unittest.main()
