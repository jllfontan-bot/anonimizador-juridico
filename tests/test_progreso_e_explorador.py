"""Melloras da interface web pedidas tras probar o instalador en Windows:

- explorador de cartafoles que amosa TAMÉN os ficheiros (en gris, non
  escollibles) -- `explorador.py` e as rutas `/explorar*`;
- traballos en segundo plano con progreso (`/traballo*`): ficheiro actual,
  porcentaxe e aviso de remate, sen tirar todo o lote se un ficheiro falla;
- ao confirmar unha revisión, os saltos de liña `\\r\\n` que envía o
  navegador NON deben facer que un PDF/.docx caia a texto plano.
"""
import os

# Estas probas escribíronse coa interface e os códigos en galego (PERSOA_...).
os.environ.setdefault("ANONIMIZADOR_IDIOMA", "gl")

import sys
import tempfile
import time
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from anonimizador import desktop, explorador, web


class TestExplorador(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)

    def tearDown(self):
        self._tmp.cleanup()

    def test_amosa_cartafoles_primeiro_e_tamen_os_ficheiros(self):
        (self.dir / "zeta").mkdir()
        (self.dir / "Alfa").mkdir()
        (self.dir / "b_escrito.docx").write_text("x")
        (self.dir / "A_demanda.pdf").write_text("x")
        (self.dir / "plano.dwg").write_text("x")
        (self.dir / ".oculto").write_text("x")

        r = explorador.listar(str(self.dir))

        self.assertEqual(r["ruta"], str(self.dir))
        self.assertEqual(
            [(e["nome"], e["tipo"]) for e in r["entradas"]],
            [("Alfa", "cartafol"), ("zeta", "cartafol"),
             ("A_demanda.pdf", "ficheiro"), ("b_escrito.docx", "ficheiro"), ("plano.dwg", "ficheiro")],
        )
        soportado = {e["nome"]: e["soportado"] for e in r["entradas"] if e["tipo"] == "ficheiro"}
        self.assertEqual(soportado, {"A_demanda.pdf": True, "b_escrito.docx": True, "plano.dwg": False})
        self.assertEqual((r["n_cartafoles"], r["n_soportados"]), (2, 2))
        self.assertIsNone(r["erro"])

    def test_pai_permite_subir(self):
        (self.dir / "sub").mkdir()
        r = explorador.listar(str(self.dir / "sub"))
        self.assertEqual(r["pai"], str(self.dir))

    def test_ruta_que_e_un_ficheiro_lista_o_seu_cartafol(self):
        (self.dir / "a.pdf").write_text("x")
        r = explorador.listar(str(self.dir / "a.pdf"))
        self.assertEqual(r["ruta"], str(self.dir))

    def test_ruta_inexistente_sube_ao_primeiro_pai_que_existe(self):
        r = explorador.listar(str(self.dir / "non" / "existe" / "aqui"))
        self.assertEqual(r["ruta"], str(self.dir))

    def test_ruta_relativa_resolvese_a_absoluta(self):
        r = explorador.listar("./datos_rexistro_que_non_existe")
        self.assertTrue(Path(r["ruta"]).is_absolute())

    def test_sen_ruta_empeza_nun_cartafol_de_usuario(self):
        r = explorador.listar(None)
        self.assertTrue(Path(r["ruta"]).is_dir())

    def test_atallos_inclúen_sempre_a_vista_de_equipo(self):
        r = explorador.listar(str(self.dir))
        self.assertEqual(r["atallos"][-1]["ruta"], explorador.EQUIPO)

    def test_vista_de_equipo_lista_a_raiz_en_mac_e_linux(self):
        with mock.patch.object(explorador.sys, "platform", "linux"):
            r = explorador.listar(explorador.EQUIPO)
        self.assertEqual(r["ruta"], "/")
        self.assertIsNone(r["pai"])

    def test_vista_de_equipo_en_windows_lista_as_unidades(self):
        with mock.patch.object(explorador.sys, "platform", "win32"), \
             mock.patch.object(explorador, "_unidades_windows", return_value=["C:\\", "V:\\"]):
            r = explorador.listar(explorador.EQUIPO)
        self.assertEqual(r["ruta"], "")
        self.assertEqual([e["nome"] for e in r["entradas"]], ["C:\\", "V:\\"])
        self.assertTrue(all(e["tipo"] == "cartafol" for e in r["entradas"]))

    def test_desde_a_raiz_dunha_unidade_windows_sube_a_vista_de_equipo(self):
        with mock.patch.object(explorador.sys, "platform", "win32"):
            self.assertEqual(explorador._pai(Path("/")), explorador.EQUIPO)

    def test_erro_de_permisos_devolve_mensaxe_en_vez_de_romper(self):
        with mock.patch.object(explorador.os, "scandir", side_effect=PermissionError(13, "Permission denied")):
            r = explorador.listar(str(self.dir))
        self.assertIn("Non se pode ler", r["erro"])
        self.assertEqual(r["entradas"], [])

    def test_crear_cartafol(self):
        novo = explorador.crear_cartafol(str(self.dir), "  Saídas anonimizadas ")
        self.assertTrue(Path(novo).is_dir())
        self.assertEqual(Path(novo).name, "Saídas anonimizadas")

    def test_crear_cartafol_rexeita_nomes_invalidos_e_duplicados(self):
        (self.dir / "existe").mkdir()
        for nome in ("", "  ", "..", "a/b", "a\\b", "co:lon", "estrela*", "punto."):
            with self.subTest(nome=nome), self.assertRaises(ValueError):
                explorador.crear_cartafol(str(self.dir), nome)
        with self.assertRaises(ValueError):
            explorador.crear_cartafol(str(self.dir), "existe")
        with self.assertRaises(ValueError):
            explorador.crear_cartafol(str(self.dir / "non_existe"), "x")


class TestRutasExplorador(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.client = web.app.test_client()

    def tearDown(self):
        self._tmp.cleanup()

    def test_explorar_devolve_json_con_ficheiros(self):
        (self.dir / "a.pdf").write_text("x")
        datos = self.client.get("/explorar", query_string={"ruta": str(self.dir)}).get_json()
        self.assertEqual([e["nome"] for e in datos["entradas"]], ["a.pdf"])

    def test_nova_carpeta(self):
        r = self.client.post("/explorar/nova-carpeta", data={"pai": str(self.dir), "nome": "saidas"})
        self.assertEqual(r.status_code, 200)
        self.assertTrue((self.dir / "saidas").is_dir())

    def test_nova_carpeta_con_nome_invalido_da_400_cun_texto(self):
        r = self.client.post("/explorar/nova-carpeta", data={"pai": str(self.dir), "nome": "a/b"})
        self.assertEqual(r.status_code, 400)
        self.assertIn("non válido", r.get_json()["erro"])


class _TraballoBase(unittest.TestCase):
    def setUp(self):
        web.SESIONS.clear()
        web.TRABALLOS.clear()
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.rexistro_dir = str(self.dir / "rexistro")
        self.client = web.app.test_client()

    def tearDown(self):
        self._tmp.cleanup()

    def _iniciar(self, ruta, modo="automatico", **extra):
        return self.client.post(
            "/traballo", data={"ruta": str(ruta), "rexistro": self.rexistro_dir, "modo": modo, **extra}
        )

    def _agardar(self, jid, tempo_maximo=180):
        fin = time.time() + tempo_maximo
        while time.time() < fin:
            estado = self.client.get(f"/traballo/{jid}/estado").get_json()
            if estado["estado"] != "en_curso":
                return estado
            time.sleep(0.05)
        self.fail("o traballo non rematou a tempo")


class TestTraballosConProgreso(_TraballoBase):
    def _lote(self, n=3):
        cartafol = self.dir / "lote"
        cartafol.mkdir()
        for i in range(n):
            (cartafol / f"doc{i}.txt").write_text(f"Don Antonio Pérez García, DNI 1234567{i}Z.", encoding="utf-8")
        return cartafol

    def test_lote_automatico_informa_de_progreso_e_remata(self):
        cartafol = self._lote(3)
        r = self._iniciar(cartafol)
        self.assertEqual(r.status_code, 200)
        datos = r.get_json()
        self.assertEqual(datos["total"], 3)

        estado = self._agardar(datos["id"])
        self.assertEqual(estado["estado"], "rematado")
        self.assertEqual((estado["feitos"], estado["total"], estado["porcentaxe"]), (3, 3, 100))
        self.assertIsNone(estado["actual"])

        pagina = self.client.get(f"/traballo/{datos['id']}/resultado").get_data(as_text=True)
        self.assertIn("Proceso terminado", pagina)
        self.assertIn("3 de 3 ficheiro(s) anonimizado(s) correctamente", pagina)
        for i in range(3):
            self.assertTrue((cartafol / f"doc{i}.anonimizado.txt").exists())

    def test_o_estado_ten_o_ficheiro_actual_e_porcentaxe_intermedia(self):
        cartafol = self._lote(2)
        pausa, liberar = __import__("threading").Event(), __import__("threading").Event()
        orixinal = web._procesar_automatico

        def lento(ficheiro, *args, **kwargs):
            if ficheiro.endswith("doc1.txt"):
                pausa.set()
                liberar.wait(30)
            return orixinal(ficheiro, *args, **kwargs)

        with mock.patch.object(web, "_procesar_automatico", side_effect=lento):
            jid = self._iniciar(cartafol).get_json()["id"]
            self.assertTrue(pausa.wait(60))
            estado = self.client.get(f"/traballo/{jid}/estado").get_json()
            self.assertEqual(estado["estado"], "en_curso")
            self.assertEqual(estado["actual"], "doc1.txt")
            self.assertEqual((estado["feitos"], estado["total"], estado["porcentaxe"]), (1, 2, 50))
            liberar.set()
            self.assertEqual(self._agardar(jid)["estado"], "rematado")

    def test_un_ficheiro_que_falla_non_tira_o_resto_do_lote(self):
        cartafol = self._lote(3)
        orixinal = web._procesar_automatico

        def falla_un(ficheiro, *args, **kwargs):
            if ficheiro.endswith("doc1.txt"):
                raise RuntimeError("PDF corrupto")
            return orixinal(ficheiro, *args, **kwargs)

        with mock.patch.object(web, "_procesar_automatico", side_effect=falla_un):
            jid = self._iniciar(cartafol).get_json()["id"]
            estado = self._agardar(jid)

        self.assertEqual(estado["estado"], "rematado")
        self.assertEqual((estado["feitos"], estado["n_erros"]), (3, 1))
        pagina = self.client.get(f"/traballo/{jid}/resultado").get_data(as_text=True)
        self.assertIn("2 de 3 ficheiro(s) anonimizado(s) correctamente", pagina)
        self.assertIn("1 con erro", pagina)
        self.assertIn("doc1.txt", pagina)
        self.assertIn("PDF corrupto", pagina)

    def test_ruta_inexistente_ou_cartafol_sen_documentos_da_400_cun_texto(self):
        r = self._iniciar(self.dir / "non_existe")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Non se atopa", r.get_json()["erro"])

        baleiro = self.dir / "baleiro"
        baleiro.mkdir()
        (baleiro / "plano.dwg").write_bytes(b"x")
        r = self._iniciar(baleiro)
        self.assertEqual(r.status_code, 400)
        self.assertIn("ningún documento soportado", r.get_json()["erro"])

    def test_estado_dun_traballo_descoñecido_da_404(self):
        self.assertEqual(self.client.get("/traballo/nonexiste/estado").status_code, 404)

    def test_resultado_dun_traballo_en_curso_volve_ao_inicio(self):
        cartafol = self._lote(1)
        liberar = __import__("threading").Event()
        orixinal = web._procesar_automatico
        with mock.patch.object(
            web, "_procesar_automatico",
            side_effect=lambda *a, **k: (liberar.wait(30), orixinal(*a, **k))[1],
        ):
            jid = self._iniciar(cartafol).get_json()["id"]
            r = self.client.get(f"/traballo/{jid}/resultado")
            self.assertEqual(r.status_code, 302)
            liberar.set()
            self._agardar(jid)

    def test_revisar_un_ficheiro_leva_directo_a_revision(self):
        doc = self.dir / "escrito.txt"
        doc.write_text("Don Antonio Pérez García, DNI 12345678Z.", encoding="utf-8")
        jid = self._iniciar(doc, modo="revisar").get_json()["id"]
        self._agardar(jid)
        r = self.client.get(f"/traballo/{jid}/resultado")
        self.assertEqual(r.status_code, 302)
        self.assertIn("/revisar/", r.headers["Location"])
        self.assertEqual(len(web.SESIONS), 1)

    def test_revisar_un_cartafol_crea_unha_sesion_por_ficheiro_e_volve_ao_inicio(self):
        cartafol = self._lote(3)
        jid = self._iniciar(cartafol, modo="revisar").get_json()["id"]
        self._agardar(jid)
        r = self.client.get(f"/traballo/{jid}/resultado")
        self.assertEqual(r.status_code, 302)
        self.assertIn("creados=3", r.headers["Location"])
        self.assertEqual(len(web.SESIONS), 3)

    def test_o_procesado_sincrono_de_sempre_segue_funcionando(self):
        """`/procesar` (sen JavaScript) non cambia."""
        doc = self.dir / "escrito.txt"
        doc.write_text("Don Antonio Pérez García, DNI 12345678Z.", encoding="utf-8")
        r = self.client.post(
            "/procesar", data={"ruta": str(doc), "rexistro": self.rexistro_dir, "modo": "automatico"}
        )
        self.assertIn("Anonimización automática", r.get_data(as_text=True))
        self.assertTrue((self.dir / "escrito.anonimizado.txt").exists())

    def test_traballos_rematados_vellos_esquécense(self):
        web.TRABALLOS["vello"] = {"fin": time.time() - web._VIDA_TRABALLO_REMATADO_S - 5}
        web.TRABALLOS["novo"] = {"fin": time.time()}
        web.TRABALLOS["en_curso"] = {"fin": None}
        web._esquecer_traballos_vellos()
        self.assertEqual(sorted(web.TRABALLOS), ["en_curso", "novo"])


class TestCancelarLote(_TraballoBase):
    def _lote(self, n):
        cartafol = self.dir / "lote"
        cartafol.mkdir()
        for i in range(n):
            (cartafol / f"doc{i}.txt").write_text(f"Don Antonio Pérez García, DNI 1234567{i}Z.", encoding="utf-8")
        return cartafol

    def _cancelar_durante_o_segundo(self, cartafol, modo):
        import threading

        pausa, liberar = threading.Event(), threading.Event()
        nome = "_procesar_automatico" if modo == "automatico" else "_crear_sesion"
        orixinal = getattr(web, nome)

        def lento(ficheiro, *args, **kwargs):
            if ficheiro.endswith("doc1.txt"):
                pausa.set()
                liberar.wait(30)
            return orixinal(ficheiro, *args, **kwargs)

        with mock.patch.object(web, nome, side_effect=lento):
            jid = self._iniciar(cartafol, modo=modo).get_json()["id"]
            self.assertTrue(pausa.wait(60))
            r = self.client.post(f"/traballo/{jid}/cancelar")
            self.assertTrue(r.get_json()["cancelando"])
            self.assertEqual(self.client.get(f"/traballo/{jid}/estado").get_json()["estado"], "en_curso")
            liberar.set()
            return jid, self._agardar(jid)

    def test_cancelar_remata_o_ficheiro_actual_e_non_empeza_o_seguinte(self):
        cartafol = self._lote(3)
        jid, estado = self._cancelar_durante_o_segundo(cartafol, "automatico")

        self.assertEqual(estado["estado"], "cancelado")
        self.assertEqual((estado["feitos"], estado["total"], estado["porcentaxe"]), (2, 3, 66))
        self.assertTrue((cartafol / "doc0.anonimizado.txt").exists())
        self.assertTrue((cartafol / "doc1.anonimizado.txt").exists())
        self.assertFalse((cartafol / "doc2.anonimizado.txt").exists())

        pagina = self.client.get(f"/traballo/{jid}/resultado").get_data(as_text=True)
        self.assertIn("Proceso cancelado", pagina)
        self.assertIn("os 1 ficheiro(s) restantes non se tocaron", pagina)
        self.assertIn("2 de 3 ficheiro(s) anonimizado(s) correctamente", pagina)

    def test_cancelar_en_modo_revisar_deixa_as_sesions_xa_creadas(self):
        cartafol = self._lote(3)
        jid, estado = self._cancelar_durante_o_segundo(cartafol, "revisar")

        self.assertEqual(estado["estado"], "cancelado")
        self.assertEqual(len(web.SESIONS), 2)
        pagina = self.client.get(f"/traballo/{jid}/resultado").get_data(as_text=True)
        self.assertIn("Proceso cancelado: 2 de 3", pagina)
        self.assertIn("Revisións pendentes", pagina)

    def test_cancelar_cando_xa_estaba_no_ultimo_ficheiro_conta_coma_rematado(self):
        cartafol = self._lote(1)
        jid = self._iniciar(cartafol).get_json()["id"]
        self._agardar(jid)
        self.client.post(f"/traballo/{jid}/cancelar")
        self.assertEqual(self.client.get(f"/traballo/{jid}/estado").get_json()["estado"], "rematado")

    def test_cancelar_un_traballo_descoñecido_da_404(self):
        self.assertEqual(self.client.post("/traballo/nonexiste/cancelar").status_code, 404)


class TestConfirmarConSaltosDeLinhaDoNavegador(unittest.TestCase):
    """Bug real: un <textarea> envía "\\r\\n"; o separador de páxina dun PDF
    ("\\n\\f\\n") deixaba de coincidir e o PDF gardábase SEMPRE coma .txt."""

    def setUp(self):
        web.SESIONS.clear()
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.rexistro_dir = str(self.dir / "rexistro")
        self.client = web.app.test_client()

    def tearDown(self):
        self._tmp.cleanup()

    def _revisar(self, ruta):
        r = self.client.post(
            "/procesar", data={"ruta": str(ruta), "rexistro": self.rexistro_dir}, follow_redirects=True
        )
        return r.request.path.rsplit("/", 1)[-1]

    def _confirmar_como_navegador(self, sid, texto):
        texto_crlf = texto.replace("\r\n", "\n").replace("\n", "\r\n")
        return self.client.post(
            f"/revisar/{sid}/confirmar", data={"texto_final": texto_crlf, "rexistro_dir": self.rexistro_dir}
        ).get_data(as_text=True)

    def test_pdf_de_varias_paxinas_conserva_o_formato(self):
        import fitz

        pdf = self.dir / "escrito.pdf"
        d = fitz.open()
        for i in range(3):
            p = d.new_page()
            p.insert_text((72, 100), f"Pagina {i + 1}. Don Antonio Perez Garcia, DNI 12345678Z.", fontsize=11)
        d.save(str(pdf))

        sid = self._revisar(pdf)
        pagina = self._confirmar_como_navegador(sid, web.SESIONS[sid]["texto_anonimizado"])

        self.assertNotIn("Non se puido conservar o formato", pagina)
        saida = self.dir / "escrito.anonimizado.pdf"
        self.assertTrue(saida.exists())
        self.assertFalse((self.dir / "escrito.anonimizado.txt").exists())
        self.assertEqual(len(fitz.open(str(saida))), 3)

    def test_docx_non_deixa_un_retorno_de_carro_solto_en_cada_paragrafo(self):
        import docx

        ruta = self.dir / "escrito.docx"
        d = docx.Document()
        d.add_paragraph("Don Antonio Pérez García, DNI 12345678Z.")
        d.add_paragraph("Segundo parágrafo sen datos.")
        d.save(str(ruta))

        sid = self._revisar(ruta)
        esperado = web.SESIONS[sid]["texto_anonimizado"].split("\n")
        self._confirmar_como_navegador(sid, web.SESIONS[sid]["texto_anonimizado"])

        saida = docx.Document(str(self.dir / "escrito.anonimizado.docx"))
        self.assertEqual([p.text for p in saida.paragraphs], esperado)  # sen "\r" solto


class TestDniConPuntosNoPdf(unittest.TestCase):
    def test_pdf_con_dni_con_puntos_e_letra_na_liña_seguinte_queda_sen_dni(self):
        import fitz

        from anonimizador.pipeline import anonimizar_e_gardar
        from anonimizador.registry import RexistroCorrespondencias

        with tempfile.TemporaryDirectory() as tmp:
            entrada, saida = Path(tmp) / "escrito.pdf", Path(tmp) / "escrito.anonimizado.pdf"
            d = fitz.open()
            p = d.new_page()
            p.insert_text((72, 100), "Doña Ana Pérez Gómez, con DNI 12.345.678-", fontsize=11)
            p.insert_text((72, 118), "V, vecina de Lugo, y don Luis Roel Salcedo con DNI 87.654.321-G.", fontsize=11)
            d.save(str(entrada))
            anonimizar_e_gardar(str(entrada), str(saida), RexistroCorrespondencias(str(Path(tmp) / "reg")))
            texto = fitz.open(str(saida))[0].get_text()
        self.assertNotRegex(texto, r"\d{2}\.\d{3}\.\d{3}")
        self.assertNotIn("12.345.678", texto)
        self.assertNotIn("87.654.321", texto)


class TestDesktopSaidasEXanela(unittest.TestCase):
    def test_as_ligazons_de_xanela_nova_non_saen_ao_navegador(self):
        import webview

        self.assertFalse(webview.settings["OPEN_EXTERNAL_LINKS_IN_BROWSER"])

    def test_sen_consola_as_saidas_van_a_un_log_no_cartafol_de_traballo(self):
        import os

        tmp = tempfile.TemporaryDirectory()
        anterior = os.getcwd()
        os.chdir(tmp.name)
        try:
            with mock.patch.object(sys, "stdout", None), mock.patch.object(sys, "stderr", None):
                ruta = desktop._redirixir_saidas_se_non_hai_consola()
                print("hola desde pythonw")
                self.assertIs(sys.stdout, sys.stderr)
                sys.stdout.close()
            self.assertEqual(ruta.name, "anonimizador.log")
            contido = ruta.read_text(encoding="utf-8")
            self.assertIn("arranque", contido)
            self.assertIn("hola desde pythonw", contido)
        finally:
            os.chdir(anterior)
            tmp.cleanup()

    def test_o_log_escribese_aínda_que_pywebview_substitúa_stdout_ao_importarse(self):
        """Bug real en Windows: `import webview` cambia `sys.stdout=None` por
        `os.devnull`, así que había que redirixir ANTES de importalo. Proba
        cun intérprete novo (sen consola simulada) e comproba o ficheiro."""
        import subprocess

        with tempfile.TemporaryDirectory() as tmp:
            codigo = (
                "import sys; sys.stdout = None; sys.stderr = None\n"
                "import anonimizador.desktop\n"
                "print('despois de importar')\n"
            )
            resultado = subprocess.run(
                [sys.executable, "-c", codigo], cwd=tmp, capture_output=True, text=True,
                env={**__import__("os").environ, "PYTHONPATH": str(Path(__file__).resolve().parents[1] / "src")},
            )
            log = Path(tmp) / "anonimizador.log"
            self.assertTrue(log.exists(), resultado.stderr)
            self.assertIn("despois de importar", log.read_text(encoding="utf-8"))

    def test_con_consola_non_se_toca_nada(self):
        self.assertIsNone(desktop._redirixir_saidas_se_non_hai_consola())

    def test_en_windows_pregunta_antes_de_abrir_o_navegador(self):
        with mock.patch.object(desktop.sys, "platform", "win32"), \
             mock.patch("ctypes.windll", create=True) as windll:
            windll.user32.MessageBoxW.return_value = 7  # IDNO
            self.assertFalse(desktop._preguntar_abrir_no_navegador(RuntimeError("sen WebView2")))
            windll.user32.MessageBoxW.return_value = 6  # IDYES
            self.assertTrue(desktop._preguntar_abrir_no_navegador(RuntimeError("sen WebView2")))
            self.assertIn("sen WebView2", windll.user32.MessageBoxW.call_args[0][1])

    def test_en_windows_se_o_usuario_di_que_non_o_navegador_non_se_abre(self):
        with mock.patch.object(desktop, "webview") as webview_falso, \
             mock.patch.object(desktop.web, "executar"), \
             mock.patch.object(desktop, "_agardar_servidor"), \
             mock.patch.object(desktop, "_preguntar_abrir_no_navegador", return_value=False), \
             mock.patch.object(desktop, "webbrowser") as navegador_falso:
            webview_falso.start.side_effect = RuntimeError("sen WebView2")
            desktop.executar(porto=5126)
            navegador_falso.open.assert_not_called()

    def test_sen_webview2_en_windows_pregunta_en_vez_de_abrir_unha_xanela_rota(self):
        """pywebview NON falla sen WebView2 (cae en silencio a Internet
        Explorer): hai que comprobalo antes de crear a xanela."""
        with mock.patch.object(desktop, "webview") as webview_falso, \
             mock.patch.object(desktop.web, "executar"), \
             mock.patch.object(desktop, "_agardar_servidor"), \
             mock.patch.object(desktop, "_webview2_instalado", return_value=False), \
             mock.patch.object(desktop, "_preguntar_abrir_no_navegador", return_value=False) as pregunta, \
             mock.patch.object(desktop, "webbrowser") as navegador_falso:
            desktop.executar(porto=5127)
            webview_falso.create_window.assert_not_called()
            navegador_falso.open.assert_not_called()
            self.assertIn("WebView2", str(pregunta.call_args[0][0]))

    def test_sen_webview2_e_respondendo_que_si_abre_o_navegador(self):
        with mock.patch.object(desktop, "webview") as webview_falso, \
             mock.patch.object(desktop.web, "executar"), \
             mock.patch.object(desktop, "_agardar_servidor"), \
             mock.patch.object(desktop, "_webview2_instalado", return_value=False), \
             mock.patch.object(desktop, "_preguntar_abrir_no_navegador", return_value=True), \
             mock.patch.object(desktop, "webbrowser") as navegador_falso:
            desktop.executar(porto=5128)
            webview_falso.create_window.assert_not_called()
            navegador_falso.open.assert_called_once_with("http://127.0.0.1:5128/")

    def test_webview2_fora_de_windows_conta_sempre_coma_dispoñible(self):
        with mock.patch.object(desktop.sys, "platform", "darwin"):
            self.assertTrue(desktop._webview2_instalado())

    def test_relanzamento_da_barra_de_tareas_usa_o_lanzador_vbs_da_instalacion(self):
        """Bug real: anclar a app aberta e pechala deixaba a icona sen función
        (Windows ancla `pythonw.exe` sen argumentos)."""
        with tempfile.TemporaryDirectory() as tmp:
            raiz = Path(tmp)
            falso = raiz / "anonimizador-xuridico" / "src" / "anonimizador" / "desktop.py"
            with mock.patch.object(desktop, "__file__", str(falso)):
                self.assertIsNone(desktop._comando_de_relanzamento())  # sen lanzador: desenvolvemento
                (raiz / "Iniciar-Anonimizador.vbs").write_text("x")
                (raiz / "anonimizador.ico").write_bytes(b"x")
                comando, icona = desktop._comando_de_relanzamento()
        self.assertIn("wscript.exe", comando)
        self.assertIn("Iniciar-Anonimizador.vbs", comando)
        self.assertTrue(icona.endswith("anonimizador.ico,0"))

    def test_configurar_boton_da_barra_garda_identidade_e_comando_de_relanzamento(self):
        almacen = mock.Mock()
        propsys = mock.Mock(**{"SHGetPropertyStoreForWindow.return_value": almacen})
        propsys.PROPVARIANTType.side_effect = lambda v: ("VT", v)
        pscon = mock.Mock(
            PKEY_AppUserModel_ID="ID", PKEY_AppUserModel_RelaunchCommand="CMD",
            PKEY_AppUserModel_RelaunchDisplayNameResource="NOME", PKEY_AppUserModel_RelaunchIconResource="ICONA",
        )
        modulos = {"pythoncom": mock.Mock(), "win32com": mock.Mock(propsys=propsys), "win32com.propsys": mock.Mock(propsys=propsys, pscon=pscon)}
        with mock.patch.object(desktop.sys, "platform", "win32"), mock.patch.dict(sys.modules, modulos), \
             mock.patch.object(desktop, "_comando_de_relanzamento", return_value=('"wscript.exe" "x.vbs"', "a.ico,0")):
            self.assertTrue(desktop._configurar_boton_da_barra(1234))
        gardado = {c.args[0]: c.args[1][1] for c in almacen.SetValue.call_args_list}
        self.assertEqual(gardado["ID"], desktop._AUMID)
        self.assertEqual(gardado["CMD"], '"wscript.exe" "x.vbs"')
        self.assertEqual(gardado["ICONA"], "a.ico,0")
        almacen.Commit.assert_called_once()

    def test_configurar_boton_da_barra_nunca_lanza_e_fora_de_windows_non_fai_nada(self):
        with mock.patch.object(desktop.sys, "platform", "darwin"):
            self.assertFalse(desktop._configurar_boton_da_barra(1))
        with mock.patch.object(desktop.sys, "platform", "win32"), mock.patch.dict(sys.modules, {"pythoncom": None}):
            self.assertFalse(desktop._configurar_boton_da_barra(1))  # ImportError controlado

    def test_fóra_de_windows_a_reserva_abre_o_navegador_sen_preguntar(self):
        with mock.patch.object(desktop.sys, "platform", "linux"):
            self.assertTrue(desktop._preguntar_abrir_no_navegador(RuntimeError("sen GTK")))


if __name__ == "__main__":
    unittest.main()
