import os

# Estas probas escribíronse coa interface e os códigos en galego (PERSOA_...).
os.environ.setdefault("ANONIMIZADOR_IDIOMA", "gl")
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from anonimizador.pipeline import anonimizar_zip
from anonimizador.registry import RexistroCorrespondencias


def _crear_zip_de_proba(ruta: Path) -> None:
    with tempfile.TemporaryDirectory() as d:
        docx_path = Path(d) / "escrito.docx"
        doc = docx.Document()
        r = doc.add_paragraph().add_run("Manuel Ferrándiz Salcedo, con DNI 30014811-Z, declarou.")
        r.bold = True
        doc.save(str(docx_path))

        with zipfile.ZipFile(ruta, "w") as zf:
            zf.writestr(
                "carpeta/doc1.txt",
                "Manuel Ferrándiz Salcedo, con DNI 30014811-Z, presentou a demanda.",
            )
            zf.writestr(
                "carpeta/doc2.txt",
                "Segundo consta, Manuel Ferrándiz Salcedo (DNI 30014811-Z) reside en Cangas.",
            )
            zf.write(str(docx_path), arcname="carpeta/escrito.docx")
            zf.writestr("planos/plano1.dwg", b"\x00\x01binario-non-analizable")


class TestAnonimizarZip(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.rexistro = RexistroCorrespondencias(str(self.dir / "rexistro"))

    def tearDown(self):
        self._tmp.cleanup()

    def test_procesa_txt_e_docx_e_omite_non_soportados(self):
        entrada = self.dir / "expediente.zip"
        _crear_zip_de_proba(entrada)
        saida = self.dir / "expediente.anonimizado.zip"

        resultado = anonimizar_zip(str(entrada), str(saida), self.rexistro)

        self.assertIn("carpeta/doc1.txt", resultado["ficheiros"])
        self.assertIn("carpeta/doc2.txt", resultado["ficheiros"])
        self.assertIn("carpeta/escrito.docx", resultado["ficheiros"])
        self.assertEqual(resultado["omitidos"], ["planos/plano1.dwg"])

        with zipfile.ZipFile(saida) as zf:
            nomes = zf.namelist()
            self.assertIn("carpeta/doc1.anonimizado.txt", nomes)
            self.assertIn("carpeta/doc2.anonimizado.txt", nomes)
            self.assertIn("carpeta/escrito.anonimizado.docx", nomes)
            # o binario non analizable NUNCA se copia ao zip de saida
            self.assertFalse(any("plano1" in n for n in nomes))

            texto1 = zf.read("carpeta/doc1.anonimizado.txt").decode("utf-8")
            texto2 = zf.read("carpeta/doc2.anonimizado.txt").decode("utf-8")
            self.assertNotIn("Manuel Ferrándiz Salcedo", texto1)
            self.assertNotIn("Manuel Ferrándiz Salcedo", texto2)

    def test_mesma_persoa_mesmo_codigo_entre_ficheiros_do_zip(self):
        """A consistencia por DNI ten que funcionar tamén entre ficheiros
        DISTINTOS dun mesmo .zip, non só entre documentos soltos."""
        entrada = self.dir / "expediente.zip"
        _crear_zip_de_proba(entrada)
        saida = self.dir / "expediente.anonimizado.zip"

        resultado = anonimizar_zip(str(entrada), str(saida), self.rexistro)

        codigo1 = resultado["ficheiros"]["carpeta/doc1.txt"]["persoas"][0]["codigo"]
        codigo2 = resultado["ficheiros"]["carpeta/doc2.txt"]["persoas"][0]["codigo"]
        codigo3 = resultado["ficheiros"]["carpeta/escrito.docx"]["persoas"][0]["codigo"]
        self.assertEqual(codigo1, codigo2)
        self.assertEqual(codigo1, codigo3)

    def test_docx_dentro_do_zip_conserva_formato(self):
        entrada = self.dir / "expediente.zip"
        _crear_zip_de_proba(entrada)
        saida = self.dir / "expediente.anonimizado.zip"

        anonimizar_zip(str(entrada), str(saida), self.rexistro)

        with zipfile.ZipFile(saida) as zf:
            with tempfile.TemporaryDirectory() as d:
                ruta_docx = Path(d) / "escrito.docx"
                ruta_docx.write_bytes(zf.read("carpeta/escrito.anonimizado.docx"))
                doc = docx.Document(str(ruta_docx))
                p = doc.paragraphs[0]
                self.assertIn("PERSOA_", p.text)
                self.assertTrue(p.runs[0].bold)


if __name__ == "__main__":
    unittest.main()
