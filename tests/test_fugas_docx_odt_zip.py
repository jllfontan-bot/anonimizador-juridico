"""Fugas críticas de datos persoais atopadas en .docx/.odt/.zip (informe
externo): cabeceiras/pés de páxina e táboas aniñadas en .docx (A2),
hyperlinks que descolocaban as substitucións (A7), control de cambios --
insercións non detectadas e delecións/autores que quedaban no XML (A8),
metadatos do paquete (.docx `docProps/*.xml`, .odt `meta.xml`, A3),
títulos/cabeceiras/anotacións en .odt (A9) e nomes de cartafol dentro dun
.zip, ademais de lixo de macOS (A10).

Cada test constrúe un documento SINTÉTICO (datos de proba, nunca reais),
anonimízao co pipeline completo e comproba a SAÍDA abrindo o resultado
coma un .zip e buscando o dato orixinal en TODOS os XML internos -- non
abonda con mirar só o texto que devolve python-docx/odfpy, porque
xustamente aí é onde se escondían estas fugas."""
import os

# Estas probas escribíronse coa interface e os códigos en galego (PERSOA_...).
os.environ.setdefault("ANONIMIZADOR_IDIOMA", "gl")

import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

import docx
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml.ns import qn
from docx.oxml.parser import OxmlElement
from odf.dc import Creator as CreatorOdt
from odf.dc import Title as TituloMetaOdt
from odf.meta import InitialCreator as InitialCreatorOdt
from odf.office import Annotation as AnnotationOdt
from odf.opendocument import OpenDocumentText
from odf.style import Footer, Header, MasterPage, PageLayout
from odf.text import H, P

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from anonimizador.pipeline import anonimizar_e_gardar, anonimizar_zip
from anonimizador.registry import RexistroCorrespondencias


# -- axudas de construción de documentos sintéticos -------------------------

def _engadir_hyperlink(paragrafo, texto: str, url: str):
    """python-docx non ten API pública para engadir un `w:hyperlink` --
    constrúese o XML á man, coma faría Word."""
    rid = paragrafo.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = texto
    r.append(t)
    hyperlink.append(r)
    paragrafo._p.append(hyperlink)


def _engadir_insercion(paragrafo, texto: str, autor: str = "Autor Cambios"):
    """Simula texto INSERIDO con control de cambios activado (`w:ins`) --
    python-docx non ten API para isto, constrúese o XML á man."""
    ins = OxmlElement("w:ins")
    ins.set(qn("w:id"), "100")
    ins.set(qn("w:author"), autor)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = texto
    r.append(t)
    ins.append(r)
    paragrafo._p.append(ins)


def _engadir_delecion(paragrafo, texto: str, autor: str = "Autor Cambios"):
    """Simula texto BORRADO con control de cambios activado (`w:del`/
    `w:delText`) -- o texto queda no XML aínda que Word xa non o mostre."""
    delecion = OxmlElement("w:del")
    delecion.set(qn("w:id"), "200")
    delecion.set(qn("w:author"), autor)
    r = OxmlElement("w:r")
    delt = OxmlElement("w:delText")
    delt.text = texto
    r.append(delt)
    delecion.append(r)
    paragrafo._p.append(delecion)


def _engadir_move_from(paragrafo, texto: str, autor: str = "Autor Cambios"):
    """Simula a ORIXE dun "mover texto" rastrexado (`w:moveFrom`) --
    igual có `w:del`, o texto queda coma `w:delText`."""
    move_from = OxmlElement("w:moveFrom")
    move_from.set(qn("w:id"), "300")
    move_from.set(qn("w:author"), autor)
    r = OxmlElement("w:r")
    delt = OxmlElement("w:delText")
    delt.text = texto
    r.append(delt)
    move_from.append(r)
    paragrafo._p.append(move_from)


def _engadir_move_to(paragrafo, texto: str, autor: str = "Autor Cambios"):
    """Simula o DESTINO dun "mover texto" rastrexado (`w:moveTo`) --
    igual có `w:ins`, o texto vai coma `w:t` normal."""
    move_to = OxmlElement("w:moveTo")
    move_to.set(qn("w:id"), "301")
    move_to.set(qn("w:author"), autor)
    r = OxmlElement("w:r")
    t = OxmlElement("w:t")
    t.text = texto
    r.append(t)
    move_to.append(r)
    paragrafo._p.append(move_to)


def _todo_o_xml_interno(ruta) -> str:
    """Concatena o contido de TODOS os XML dentro do .docx/.odt/.zip de
    saída -- as fugas verificadas escondíanse en partes que non son o
    corpo principal (cabeceiras, comentarios, metadatos...), así que hai
    que mirar en TODAS, non só en `word/document.xml`/`content.xml`."""
    anacos = []
    with zipfile.ZipFile(ruta) as zf:
        for nome in zf.namelist():
            if nome.endswith(".xml") or nome.endswith(".rels"):
                anacos.append(zf.read(nome).decode("utf-8", errors="ignore"))
    return "\n".join(anacos)


class TestFugasDocx(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.rexistro = RexistroCorrespondencias(str(self.dir / "rexistro"))

    def tearDown(self):
        self._tmp.cleanup()

    def _anonimizar(self, doc: docx.Document, nome: str = "escrito.docx") -> Path:
        entrada = self.dir / nome
        doc.save(str(entrada))
        saida = self.dir / (Path(nome).stem + ".anonimizado.docx")
        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)
        return saida

    def test_cabeceira_e_pe_de_paxina_non_filtran_datos(self):
        """A2: cabeceira e pé de páxina (sección por defecto) NON se
        procesaban -- o nome e o DNI do despacho, e o teléfono do pé,
        quedaban intactos en header1.xml/footer1.xml."""
        doc = docx.Document()
        cabeceira = doc.sections[0].header
        cabeceira.paragraphs[0].add_run(
            "Xosé Manuel Carballo Portela, DNI 30014811Z"
        )
        pe = doc.sections[0].footer
        pe.paragraphs[0].add_run("Tel. 666123456")
        doc.add_paragraph("Corpo do escrito, sen máis datos relevantes aquí.")

        saida = self._anonimizar(doc)

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Carballo Portela", xml)
        self.assertNotIn("30014811Z", xml)
        self.assertNotIn("666123456", xml)

        doc_final = docx.Document(str(saida))
        self.assertIn("PERSOA_", doc_final.sections[0].header.paragraphs[0].text)
        self.assertIn("[TELEFONO]", doc_final.sections[0].footer.paragraphs[0].text)

    def test_cabeceira_primeira_paxina_taboa_e_celdas_combinadas(self):
        """A2: cabeceira de PRIMEIRA PÁXINA distinta da normal, unha
        TÁBOA dentro da cabeceira, e celdas COMBINADAS nesa táboa (que
        `row.cells` repite coma o mesmo obxecto -- ten que deduplicarse
        para non duplicar parágrafos nin facer fallar a escritura)."""
        doc = docx.Document()
        seccion = doc.sections[0]
        seccion.different_first_page_header_footer = True

        cabeceira_normal = seccion.header
        cabeceira_normal.paragraphs[0].add_run("Cabeceira normal: Xosé Manuel Carballo Portela")

        cabeceira_primeira = seccion.first_page_header
        cabeceira_primeira.paragraphs[0].add_run(
            "Primeira páxina: contacte con Xosé Manuel Carballo Portela"
        )
        taboa = cabeceira_primeira.add_table(rows=1, cols=2, width=docx.shared.Inches(4))
        taboa.cell(0, 0).merge(taboa.cell(0, 1))
        taboa.cell(0, 0).paragraphs[0].add_run("Responsable: Xosé Manuel Carballo Portela")

        doc.add_paragraph("Corpo do escrito, sen máis datos relevantes aquí.")

        saida = self._anonimizar(doc, "cabeceira_primeira.docx")

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Carballo Portela", xml)

    def test_hyperlink_no_medio_do_paragrafo_non_descoloca_substitucions(self):
        """A7: un `w:hyperlink` no medio do parágrafo -- `paragraph.text`
        inclúe o seu texto pero `paragraph.runs` NON, así que as
        substitucións por tramos se descolocaban (o email do hyperlink e
        o nome despois del quedaban parcialmente sen anonimizar/
        cortados mal)."""
        doc = docx.Document()
        p = doc.add_paragraph("Pode escribir a ")
        _engadir_hyperlink(p, "xmcastro@example.com", "mailto:xmcastro@example.com")
        p.add_run(" e asina Ana Belén Sotelo Revilla.")

        saida = self._anonimizar(doc, "hyperlink.docx")

        # O texto VISIBLE (word/document.xml) é o que estaba descolocado
        # -- iso é o que arranxa A7. O enderezo de destino do hyperlink
        # (`Target="mailto:..."` en word/_rels/document.xml.rels)
        # compróbase á parte máis abaixo (`test_mailto_no_destino_do_
        # hyperlink_baleirase`) -- era un límite pendente, agora arranxado.
        with zipfile.ZipFile(saida) as zf:
            documento_xml = zf.read("word/document.xml").decode("utf-8")
        self.assertNotIn("xmcastro@example.com", documento_xml)
        self.assertNotIn("Ana Belén Sotelo Revilla", documento_xml)
        self.assertNotIn("Belén Soto", documento_xml)

        doc_final = docx.Document(str(saida))
        texto = doc_final.paragraphs[0].text
        self.assertIn("[EMAIL]", texto)
        self.assertIn("PERSOA_", texto)

    def test_mailto_no_destino_do_hyperlink_baleirase(self):
        """Gravidade Media (informe do 21/09/2026): o TEXTO visible do
        hyperlink xa se anonimizaba (A7, test de arriba), pero o seu
        DESTINO real -- `Target="mailto:..."` na relación, en
        `word/_rels/document.xml.rels`, fóra de calquera parágrafo --
        quedaba intacto aínda que a palabra mostrada fose xenérica coma
        "Contacto"."""
        doc = docx.Document()
        p = doc.add_paragraph("Pode escribir a ")
        _engadir_hyperlink(
            p, "Contacto",
            "mailto:xmcastro@example.com?subject=DNI%2012345678Z",
        )
        p.add_run(".")

        saida = self._anonimizar(doc, "mailto.docx")

        with zipfile.ZipFile(saida) as zf:
            rels = zf.read("word/_rels/document.xml.rels").decode("utf-8")
        self.assertNotIn("xmcastro@example.com", rels)
        self.assertNotIn("12345678Z", rels)

    def test_texto_inserido_con_control_de_cambios_detectase(self):
        """A8: texto INSERIDO con control de cambios (`w:ins`) non
        chegaba nin a `paragraph.text` -- quedaba fóra da detección por
        completo, sen anonimizar."""
        doc = docx.Document()
        p = doc.add_paragraph("Comparece tamén ")
        _engadir_insercion(p, "Brais Lorenzo Ogando Fandiño")
        p.add_run(", maior de idade.")

        saida = self._anonimizar(doc, "insercion.docx")

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Brais Lorenzo Ogando Fandiño", xml)

        doc_final = docx.Document(str(saida))
        self.assertIn("PERSOA_", doc_final.paragraphs[0].text)

    def test_texto_borrado_e_autor_do_cambio_desaparecen(self):
        """A8: texto BORRADO con control de cambios (`w:del`/`w:delText`)
        queda no XML aínda que Word xa non o amose -- e o autor do
        cambio (`w:author`) tamén pode ser un dato persoal."""
        doc = docx.Document()
        p = doc.add_paragraph("Texto orixinal. ")
        _engadir_delecion(p, "DNI secreto do cliente: 30014811Z", autor="Leticia Viana Author")
        p.add_run("Resto do parágrafo.")

        saida = self._anonimizar(doc, "delecion.docx")

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("DNI secreto do cliente", xml)
        self.assertNotIn("30014811Z", xml)
        self.assertNotIn("w:del", xml)
        self.assertNotIn("Leticia Viana Author", xml)

        doc_final = docx.Document(str(saida))
        self.assertIn("Texto orixinal.", doc_final.paragraphs[0].text)
        self.assertIn("Resto do parágrafo.", doc_final.paragraphs[0].text)

    def test_texto_movido_detectase_no_destino_e_desaparece_na_orixe(self):
        """Gravidade Media (informe do 21/09/2026): un "mover texto"
        rastrexado (`w:moveFrom`/`w:moveTo`, distinto de `w:ins`/`w:del`)
        non se tocaba -- o DESTINO (`w:moveTo`, coma un `w:ins`) quedaba
        invisible para `paragraph.text` e escapaba sen detectar; a ORIXE
        (`w:moveFrom`, coma un `w:del`) debería desaparecer igual có
        `w:del` ao aceptar o cambio, e non o facía."""
        doc = docx.Document()
        p1 = doc.add_paragraph("Primeiro parágrafo. ")
        _engadir_move_from(p1, "Comparece Brais Lorenzo Ogando Fandiño.", autor="Notaria Movedora")
        p2 = doc.add_paragraph("Terceiro parágrafo. ")
        _engadir_move_to(p2, "Comparece Brais Lorenzo Ogando Fandiño.", autor="Notaria Movedora")

        saida = self._anonimizar(doc, "move.docx")

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Brais Lorenzo Ogando Fandiño", xml)
        self.assertNotIn("Notaria Movedora", xml)
        self.assertNotIn("w:moveFrom", xml)
        self.assertNotIn("w:moveTo", xml)

        doc_final = docx.Document(str(saida))
        self.assertNotIn("Comparece", doc_final.paragraphs[0].text)
        self.assertIn("PERSOA_", doc_final.paragraphs[1].text)

    def test_comentarios_de_revision_baleiranse(self):
        """A8: os comentarios de revisión (`word/comments.xml`) poden
        conter notas internas cun dato persoal e o nome de quen comenta
        -- non pasaban por ningunha limpeza."""
        doc = docx.Document()
        p = doc.add_paragraph("Parágrafo cun comentario de revisión.")
        doc.add_comment(
            p.runs, text="Confirmar co cliente Xosé Manuel Carballo Portela antes de enviar.",
            author="Leticia Viana Letrada",
        )

        saida = self._anonimizar(doc, "comentario.docx")

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Carballo Portela", xml)
        self.assertNotIn("Leticia Viana Letrada", xml)

    def test_metadatos_core_e_app_baleiranse(self):
        """A3 (.docx): `docProps/core.xml` (autor, quen editou por
        última vez, título, comentarios...) e `docProps/app.xml`
        (Company/Manager, sen API en python-docx) poden levar o nome
        real de quen creou/editou o documento."""
        doc = docx.Document()
        doc.add_paragraph("Corpo do escrito, sen máis datos relevantes aquí.")
        doc.core_properties.author = "Xosé Manuel Carballo Portela"
        doc.core_properties.last_modified_by = "Xosé Manuel Carballo Portela"
        doc.core_properties.title = "Escrito de Xosé Manuel Carballo Portela"
        doc.core_properties.comments = "Revisado por Xosé Manuel Carballo Portela"

        entrada = self.dir / "metadatos.docx"
        doc.save(str(entrada))
        # python-docx non ten API para docProps/app.xml -- inxéctase Company/
        # Manager á man, coma faría Word ao gardar dende un equipo corporativo.
        with zipfile.ZipFile(entrada) as zin:
            partes = {nome: zin.read(nome) for nome in zin.namelist()}
        app_xml = partes["docProps/app.xml"].decode("utf-8").replace(
            "</Properties>",
            "<Company>Bufete Carballo Portela</Company>"
            "<Manager>Xosé Manuel Carballo Portela</Manager></Properties>",
        )
        partes["docProps/app.xml"] = app_xml.encode("utf-8")
        with zipfile.ZipFile(entrada, "w") as zout:
            for nome, datos in partes.items():
                zout.writestr(nome, datos)

        saida = self.dir / "metadatos.anonimizado.docx"
        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Carballo Portela", xml)

        doc_final = docx.Document(str(saida))
        self.assertEqual(doc_final.core_properties.author, "")
        self.assertEqual(doc_final.core_properties.last_modified_by, "")
        self.assertEqual(doc_final.core_properties.title, "")
        self.assertEqual(doc_final.core_properties.comments, "")


class TestFugasOdt(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.rexistro = RexistroCorrespondencias(str(self.dir / "rexistro"))

    def tearDown(self):
        self._tmp.cleanup()

    def _anonimizar(self, doc: OpenDocumentText, nome: str = "escrito.odt") -> Path:
        entrada = self.dir / nome
        doc.save(str(entrada))
        saida = self.dir / (Path(nome).stem + ".anonimizado.odt")
        anonimizar_e_gardar(str(entrada), str(saida), self.rexistro)
        return saida

    def test_titulo_text_h_non_filtra_datos(self):
        """A9: os títulos (`text:h`) non se procesaban -- só
        `getElementsByType(P)`, que non inclúe `text:h`."""
        doc = OpenDocumentText()
        doc.text.addElement(H(outlinelevel=1, text="Recurso de Xosé Manuel Carballo Portela"))
        doc.text.addElement(P(text="Corpo do escrito, sen máis datos relevantes aquí."))

        saida = self._anonimizar(doc, "titulo.odt")

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Carballo Portela", xml)

    def test_cabeceira_e_pe_de_paxina_odt_non_filtran(self):
        """A9 (ampliación): cabeceiras/pés de páxina .odt viven en
        `styles.xml` (páxinas mestras), non en `content.xml`. NOTA: este
        caso concreto xa funcionaba antes do arranxo -- `getElementsByType`
        (odfpy) busca en TODA a árbore do documento, incluídas as páxinas
        mestras -- pero mantense coma test de regresión, xa que agora
        pasa pola nova función `_paragrafos_odt` xunto cos títulos."""
        doc = OpenDocumentText()
        doc.text.addElement(P(text="Corpo do escrito, sen máis datos relevantes aquí."))

        masterpage = MasterPage(name="Standard", pagelayoutname="PM1")
        cabeceira = Header()
        cabeceira.addElement(P(text="Xosé Manuel Carballo Portela, DNI 30014811Z"))
        masterpage.addElement(cabeceira)
        pe = Footer()
        pe.addElement(P(text="Tel. 666123456"))
        masterpage.addElement(pe)
        doc.masterstyles.addElement(masterpage)
        doc.automaticstyles.addElement(PageLayout(name="PM1"))

        saida = self._anonimizar(doc, "cabeceira.odt")

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Carballo Portela", xml)
        self.assertNotIn("30014811Z", xml)
        self.assertNotIn("666123456", xml)

    def test_anotacion_baleirase(self):
        """A9 (ampliación): as anotacións (`office:annotation`,
        comentarios de revisión) poden levar o autor e un comentario cun
        dato persoal."""
        doc = OpenDocumentText()
        paragrafo = P()
        paragrafo.addText("Parágrafo cunha anotación de revisión.")
        anotacion = AnnotationOdt()
        anotacion.addElement(CreatorOdt(text="Leticia Viana Letrada"))
        anotacion.addElement(P(text="Confirmar dato de Xosé Manuel Carballo Portela"))
        paragrafo.addElement(anotacion)
        doc.text.addElement(paragrafo)

        saida = self._anonimizar(doc, "anotacion.odt")

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Carballo Portela", xml)
        self.assertNotIn("Leticia Viana Letrada", xml)

    def test_metadatos_meta_xml_baleiranse(self):
        """A3 (.odt): `meta.xml` (creador inicial, último autor, título)
        pode levar o nome real de quen creou/editou o documento."""
        doc = OpenDocumentText()
        doc.text.addElement(P(text="Corpo do escrito, sen máis datos relevantes aquí."))
        doc.meta.addElement(InitialCreatorOdt(text="Xosé Manuel Carballo Portela"))
        doc.meta.addElement(CreatorOdt(text="Xosé Manuel Carballo Portela"))
        doc.meta.addElement(TituloMetaOdt(text="Escrito de Xosé Manuel Carballo Portela"))

        saida = self._anonimizar(doc, "metadatos.odt")

        xml = _todo_o_xml_interno(saida)
        self.assertNotIn("Carballo Portela", xml)


class TestFugasZip(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self._tmp.name)
        self.rexistro = RexistroCorrespondencias(str(self.dir / "rexistro"))

    def tearDown(self):
        self._tmp.cleanup()

    def test_nome_de_cartafol_anonimizase(self):
        """A10: o nome de ficheiro xa se anonimizaba, pero os
        compoñentes do CARTAFOL que o contén non se tocaban."""
        entrada = self.dir / "expediente.zip"
        with zipfile.ZipFile(entrada, "w") as zf:
            zf.writestr(
                "Cliente Xosé Manuel Carballo Portela/nota.txt",
                "Contido sen máis datos persoais relevantes aquí.",
            )
        saida = self.dir / "expediente.anonimizado.zip"

        resultado = anonimizar_zip(str(entrada), str(saida), self.rexistro)
        self.assertIn("Cliente Xosé Manuel Carballo Portela/nota.txt", resultado["ficheiros"])

        with zipfile.ZipFile(saida) as zf:
            nomes = zf.namelist()
            self.assertFalse(any("Carballo Portela" in n for n in nomes))
            self.assertTrue(any(n.endswith("nota.anonimizado.txt") for n in nomes))

    def test_macosx_e_appledouble_omitense(self):
        """A10: as entradas de metadatos de macOS (`__MACOSX/`, ficheiros
        `._*`) non son documentos reais -- non se deben tentar procesar
        nin incluír na saída."""
        entrada = self.dir / "expediente.zip"
        with zipfile.ZipFile(entrada, "w") as zf:
            zf.writestr("carpeta/doc.txt", "Contido normal sen datos persoais.")
            zf.writestr("__MACOSX/carpeta/._doc.txt", b"\x00\x01lixo-appledouble")
            zf.writestr("carpeta/._doc.txt", b"\x00\x01lixo-appledouble")
        saida = self.dir / "expediente.anonimizado.zip"

        resultado = anonimizar_zip(str(entrada), str(saida), self.rexistro)

        self.assertIn("__MACOSX/carpeta/._doc.txt", resultado["omitidos"])
        self.assertIn("carpeta/._doc.txt", resultado["omitidos"])
        self.assertNotIn("carpeta/._doc.txt", resultado["ficheiros"])

        with zipfile.ZipFile(saida) as zf:
            nomes = zf.namelist()
            self.assertFalse(any("__MACOSX" in n for n in nomes))
            self.assertFalse(any(n.split("/")[-1].startswith("._") for n in nomes))


if __name__ == "__main__":
    unittest.main()
